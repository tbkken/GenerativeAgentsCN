#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""独立校验教学 PPT：页数 / 备注页数 / 零位图。
用法：
    python verify_pptx.py <pptx路径> [<pptx路径> ...]
    python verify_pptx.py --scan "<book根目录>"
"""
import sys
import zipfile
import re
from pathlib import Path

try:
    from pptx import Presentation
except ImportError:
    print("需要 python-pptx：请使用 venv 中的 python")
    sys.exit(2)


def verify(path):
    p = Path(path)
    if not p.exists():
        return f"[缺失] {p}"
    try:
        prs = Presentation(str(p))
    except Exception as e:
        return f"[无法打开] {p.name} :: {e}"

    total = len(prs.slides)
    with_notes = 0
    lens = []
    for s in prs.slides:
        try:
            t = s.notes_slide.notes_text_frame.text.strip()
        except Exception:
            t = ""
        if t:
            with_notes += 1
            lens.append(len(t))

    try:
        z = zipfile.ZipFile(str(p))
        media = [n for n in z.namelist() if "/media/" in n]
    except Exception:
        media = ["<无法读取包>"]

    ok_notes = "OK " if with_notes == total else "!! "
    ok_media = "OK " if not media else "!! "
    size_kb = p.stat().st_size / 1024
    span = f"{min(lens)}-{max(lens)}" if lens else "-"
    return (
        f"{p.parent.name}/{p.name}\n"
        f"   页数 {total} | {ok_notes}备注 {with_notes}/{total} (字数 {span})\n"
        f"   {ok_media}位图 {len(media)} 个"
        + (" :: " + ", ".join(media[:4]) if media else "")
        + f"\n   大小 {size_kb:.0f} KB"
    )


def scan(root):
    root = Path(root)
    out = []
    for f in sorted(root.rglob("*.pptx")):
        if f.name.startswith("~$"):
            continue
        out.append(str(f))
    return out


if __name__ == "__main__":
    args = sys.argv[1:]
    if not args:
        print(__doc__)
        sys.exit(1)
    targets = []
    if args[0] == "--scan":
        targets = scan(args[1])
    else:
        targets = args

    bad = 0
    for t in targets:
        line = verify(t)
        print(line)
        if "!!" in line or line.startswith("[") or "缺失" in line:
            bad += 1
    print(f"\n共 {len(targets)} 个文件，异常 {bad} 个")
